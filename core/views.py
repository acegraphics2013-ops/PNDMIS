import calendar
import csv
import io
import json
from datetime import date, datetime, timedelta

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import PasswordResetForm
from django.contrib.auth.models import User
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.http import Http404, HttpResponse, HttpResponseBadRequest, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_http_methods, require_POST

from .access import (
    MANAGER,
    STAFF,
    SUPERVISOR,
    can_archive_task,
    can_archive_training,
    can_approve_deliverable,
    can_create_staff,
    can_delete_task,
    can_delete_training,
    can_edit_staff,
    can_edit_task,
    can_edit_training,
    can_manage_config,
    can_manage_deliverable,
    can_manage_shift,
    can_manage_users,
    can_review_deliverable,
    can_submit_deliverable,
    can_toggle_staff,
    can_validate_training,
    can_view_audit,
    can_view_sensitive_staff,
    can_view_staff,
    permitted_deliverables,
    permitted_shifts,
    permitted_staff,
    permitted_tasks,
    permitted_training,
    permitted_units,
    profile,
    role,
    role_required,
    task_assignable_users,
)
from .forms import (
    BulkDeliverableForm,
    CategoryForm,
    DeliverableForm,
    ProfileSelfForm,
    ReportFilterForm,
    ReviewForm,
    ShiftForm,
    StaffRecordForm,
    SubmissionForm,
    TaskForm,
    TaskStatusForm,
    TrainingForm,
    TrainingValidationForm,
    UnitForm,
    UserCreateForm,
    UserUpdateForm,
)
from .models import (
    AuditEvent,
    Category,
    Deliverable,
    Notification,
    Profile,
    ShiftCount,
    StaffRecord,
    Task,
    Training,
    Unit,
)
from .reports import build_report, filter_form, report_filename
from .services import (
    log_audit,
    notify,
    notify_roles,
    run_reminders,
    submit_deliverable,
    transition_deliverable,
    validate_training as validate_training_service,
)

DELIVERABLE_STATUSES = [key for key, _ in Deliverable.STATUSES]
TRAINING_STATUSES = [key for key, _ in Training.STATUSES]
VALIDATION_STATUSES = [key for key, _ in Training.VALIDATION]
TASK_STATUSES = [key for key, _ in Task.STATUSES]


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------
def _json_error(message, status=400, errors=None):
    data = {'ok': False, 'error': message}
    if errors is not None:
        data['errors'] = errors
    return JsonResponse(data, status=status)


def _api_login_required(view):
    def wrapped(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return _json_error('Authentication required', 401)
        try:
            profile(request.user)
            return view(request, *args, **kwargs)
        except PermissionDenied as exc:
            return _json_error(str(exc) or 'Forbidden', 403)
        except ValidationError as exc:
            return _json_error('; '.join(exc.messages) or 'Validation failed', 400)
        except Http404:
            return _json_error('Not found or no access', 404)
    wrapped.__name__ = view.__name__
    wrapped.__doc__ = view.__doc__
    return wrapped


def _page(request, queryset, per_page=20):
    paginator = Paginator(queryset, per_page)
    return paginator.get_page(request.GET.get('page'))


def _int_or_none(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _date_or_none(value):
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        return None


def _form_errors(form):
    return form.errors.get_json_data()


def _safe_cell(value):
    text = '' if value is None else str(value)
    if text.startswith(('=', '+', '-', '@')):
        return "'" + text
    return text


def _ip(request):
    return request.META.get('REMOTE_ADDR') if request else None


def _month_bounds(year, month):
    first = date(year, month, 1)
    last = date(year, month, calendar.monthrange(year, month)[1])
    return first, last


def _parse_month(value, fallback=None):
    if value:
        try:
            parsed = datetime.strptime(value, '%Y-%m')
            return parsed.year, parsed.month
        except ValueError:
            pass
    today = fallback or timezone.localdate()
    return today.year, today.month


def _task_payload_audit(task):
    return {
        'title': task.title,
        'unit': task.unit_id,
        'category': task.category_id,
        'assignee': task.assignee_id,
        'due_date': task.due_date.isoformat() if task.due_date else None,
        'priority': task.priority,
        'status': task.status,
    }


def task_to_dict(task, user):
    p = profile(user)
    return {
        'id': task.pk,
        'name': task.title,
        'title': task.title,
        'description': task.description or '',
        'category': task.unit.name if task.unit_id else '—',
        'category_id': task.unit_id,
        'unit': task.unit.name if task.unit_id else '—',
        'unit_id': task.unit_id,
        'classification': task.category.name if task.category_id else '—',
        'classification_id': task.category_id,
        'status': task.status,
        'status_display': task.get_status_display(),
        'priority': task.priority,
        'priority_display': task.get_priority_display(),
        'due_date': task.due_date.isoformat() if task.due_date else None,
        'created': task.created_at.isoformat() if task.created_at else None,
        'updated': task.updated_at.isoformat() if getattr(task, 'updated_at', None) else None,
        'assignee': task.assignee.username if task.assignee_id else '—',
        'assignee_id': task.assignee_id,
        'created_by': task.created_by.username if task.created_by_id else '—',
        'overdue': task.due_date < timezone.localdate() and task.status not in ('completed', 'archived'),
        'status_only': p.role == STAFF,
        'can_edit': can_edit_task(user, task),
        'can_archive': can_archive_task(user, task),
        'can_delete': can_delete_task(user, task),
        'can_restore': can_archive_task(user, task) and task.status == 'archived',
    }


def _dashboard_cards(role_name, values):
    if role_name == SUPERVISOR:
        return [
            {'label': 'Active staff', 'value': values['staff_count'], 'icon': 'users', 'tone': 'blue', 'url': reverse('staff')},
            {'label': 'Pending final approval', 'value': values['pending_approval'], 'icon': 'badge-check', 'tone': 'amber', 'url': reverse('deliverables') + '?status=verified'},
            {'label': 'Overdue deliverables', 'value': values['overdue'], 'icon': 'alarm-clock', 'tone': 'red', 'url': reverse('deliverables') + '?overdue=1'},
            {'label': 'PRC licenses ≤ 90 days', 'value': values['expiring_count'], 'icon': 'id-card', 'tone': 'violet', 'url': reverse('reports') + '?kind=licenses'},
        ]
    if role_name == MANAGER:
        return [
            {'label': 'Unit staff', 'value': values['staff_count'], 'icon': 'users', 'tone': 'blue', 'url': reverse('staff')},
            {'label': 'Awaiting verification', 'value': values['pending_verification'], 'icon': 'clipboard-check', 'tone': 'amber', 'url': reverse('deliverables') + '?status=submitted'},
            {'label': 'Overdue deliverables', 'value': values['overdue'], 'icon': 'alarm-clock', 'tone': 'red', 'url': reverse('deliverables') + '?overdue=1'},
            {'label': 'Tasks due soon', 'value': values['tasks_due_soon'], 'icon': 'calendar-clock', 'tone': 'violet', 'url': reverse('tasks') + '?filter=soon'},
        ]
    return [
        {'label': 'My open deliverables', 'value': values['pending'], 'icon': 'file-clock', 'tone': 'blue', 'url': reverse('deliverables')},
        {'label': 'My overdue items', 'value': values['overdue'], 'icon': 'alarm-clock', 'tone': 'red', 'url': reverse('deliverables') + '?overdue=1'},
        {'label': 'My open tasks', 'value': values['open_tasks'], 'icon': 'list-checks', 'tone': 'amber', 'url': reverse('tasks') + '?filter=open'},
        {'label': 'Training hours', 'value': values['training_hours'], 'icon': 'graduation-cap', 'tone': 'violet', 'url': reverse('training')},
    ]


# ---------------------------------------------------------------------------
# Dashboard and page shells
# ---------------------------------------------------------------------------
@login_required
def dashboard(request):
    p = profile(request.user)
    today = timezone.localdate()
    soon = today + timedelta(days=7)
    staff = permitted_staff(request.user).select_related('user', 'unit')
    deliverables = permitted_deliverables(request.user).filter(archived=False).select_related('staff__user', 'staff__unit')
    open_deliverables = deliverables.exclude(status__in=['verified', 'approved'])
    tasks = permitted_tasks(request.user).select_related('unit', 'assignee', 'category')
    open_tasks = tasks.exclude(status__in=['completed', 'archived'])

    if p.role == SUPERVISOR:
        shift_today = ShiftCount.objects.filter(date=today).select_related('unit')
        unit_summary = Unit.objects.filter(active=True).annotate(
            staff_count=Count('staffrecord', filter=Q(staffrecord__active=True), distinct=True),
            overdue_count=Count(
                'staffrecord__deliverables',
                filter=Q(staffrecord__deliverables__due_date__lt=today)
                & ~Q(staffrecord__deliverables__status__in=['verified', 'approved'])
                & Q(staffrecord__deliverables__archived=False),
                distinct=True,
            ),
        ).order_by('name')
    elif p.role == MANAGER:
        shift_today = ShiftCount.objects.filter(date=today, unit=p.unit).select_related('unit')
        unit_summary = Unit.objects.filter(pk=p.unit).annotate(
            staff_count=Count('staffrecord', filter=Q(staffrecord__active=True), distinct=True),
            overdue_count=Count(
                'staffrecord__deliverables',
                filter=Q(staffrecord__deliverables__due_date__lt=today)
                & ~Q(staffrecord__deliverables__status__in=['verified', 'approved'])
                & Q(staffrecord__deliverables__archived=False),
                distinct=True,
            ),
        )
    else:
        shift_today = ShiftCount.objects.none()
        unit_summary = Unit.objects.none()

    on_duty_total = shift_today.aggregate(total=Sum('on_duty'))['total'] or 0
    expiring = staff.filter(
        active=True, prc_expiry__isnull=False, prc_expiry__lte=today + timedelta(days=90)
    ).order_by('prc_expiry')[:10]
    expiring_count = staff.filter(
        active=True, prc_expiry__isnull=False, prc_expiry__lte=today + timedelta(days=90)
    ).count()
    training_qs = permitted_training(request.user).filter(archived=False)
    training_hours = training_qs.aggregate(total=Sum('hours'))['total'] or 0

    pending_verification = deliverables.filter(status__in=['submitted', 'review']).count()
    pending_approval = deliverables.filter(status='verified').count()
    if p.role == SUPERVISOR:
        approval_queue = deliverables.filter(status='verified').order_by('due_date')[:8]
    elif p.role == MANAGER:
        approval_queue = deliverables.filter(status__in=['submitted', 'review']).order_by('due_date')[:8]
    else:
        approval_queue = open_deliverables.order_by('due_date')[:8]

    if p.role == SUPERVISOR:
        recent_qs = AuditEvent.objects.select_related('actor')
    elif p.role == MANAGER and p.unit_id:
        recent_qs = AuditEvent.objects.filter(
            Q(actor=request.user) | Q(actor__nursing_profile__unit_id=p.unit_id)
        ).select_related('actor')
    else:
        recent_qs = AuditEvent.objects.filter(actor=request.user).select_related('actor')

    my_staff = getattr(request.user, 'staff_record', None)
    license_days = (my_staff.prc_expiry - today).days if my_staff and my_staff.prc_expiry else None

    values = {
        'staff_count': staff.filter(active=True).count(),
        'pending': open_deliverables.count(),
        'pending_verification': pending_verification,
        'pending_approval': pending_approval,
        'overdue': open_deliverables.filter(due_date__lt=today).count(),
        'expiring_count': expiring_count,
        'tasks_due_soon': open_tasks.filter(due_date__gte=today, due_date__lte=soon).count(),
        'open_tasks': open_tasks.count(),
        'training_hours': training_hours,
    }
    cards = _dashboard_cards(p.role, values)
    return render(request, 'core/dashboard.html', {
        'role': p.get_role_display(),
        'user_role': p.role,
        'ph_now': timezone.localtime(timezone.now()),
        'today': today,
        'cards': cards,
        'values': values,
        'staff_count': values['staff_count'],
        'staff_per_unit': staff.filter(active=True).values('unit__name').annotate(count=Count('id')).order_by('unit__name'),
        'shift_today': shift_today,
        'on_duty_total': on_duty_total,
        'unit_summary': unit_summary,
        'pending': values['pending'],
        'pending_verification': pending_verification,
        'pending_approval': pending_approval,
        'overdue': values['overdue'],
        'approval_queue': approval_queue,
        'approaching': open_deliverables.filter(due_date__gte=today, due_date__lte=soon).order_by('due_date')[:10],
        'overdue_list': open_deliverables.filter(due_date__lt=today).order_by('due_date')[:10],
        'deliverables': deliverables.order_by('due_date')[:12],
        'tasks': open_tasks.order_by('due_date')[:10],
        'tasks_overdue': open_tasks.filter(due_date__lt=today).count(),
        'tasks_due_soon': values['tasks_due_soon'],
        'expiring_licenses': expiring,
        'training_hours': training_hours,
        'training_validated': training_qs.filter(validation_status='validated').count(),
        'training_pending': training_qs.filter(validation_status='pending').count(),
        'recent_activity': recent_qs.order_by('-occurred_at')[:10],
        'my_staff': my_staff,
        'license_days': license_days,
        'my_training': training_qs.select_related('staff__user').order_by('-date')[:6],
        'page_title': 'Dashboard',
        'breadcrumb': 'Home / Dashboard',
    })


@login_required
def records_page(request):
    p = profile(request.user)
    units = permitted_units(request.user).filter(active=True).order_by('name')
    categories = Category.objects.filter(kind='record', active=True).order_by('name')
    return render(request, 'core/records.html', {
        'page_title': 'Records Management',
        'breadcrumb': 'Home / Records',
        'units': units,
        'categories': categories,
        'assignees': task_assignable_users(request.user),
        'can_write': p.role in (SUPERVISOR, MANAGER),
        'user_role': p.role,
    })


# ---------------------------------------------------------------------------
# JSON API for records / notifications
# ---------------------------------------------------------------------------
@_api_login_required
def api_stats(request):
    today = timezone.localdate()
    staff = permitted_staff(request.user)
    deliverables = permitted_deliverables(request.user).filter(archived=False)
    tasks = permitted_tasks(request.user)
    training = permitted_training(request.user).filter(archived=False)
    data = {
        'ok': True,
        'total_records': tasks.count() + deliverables.count(),
        'active_records': staff.filter(active=True).count(),
        'pending_records': deliverables.exclude(status__in=['verified', 'approved']).count()
        + tasks.exclude(status__in=['completed', 'archived']).count(),
        'archived_records': tasks.filter(status='archived').count() + deliverables.filter(archived=True).count(),
        'overdue': deliverables.filter(due_date__lt=today).exclude(status__in=['verified', 'approved']).count(),
        'pending_approval': deliverables.filter(status='verified').count(),
        'training_hours': float(training.aggregate(total=Sum('hours'))['total'] or 0),
        'unread_notifications': Notification.objects.filter(user=request.user, is_read=False).count(),
    }
    return JsonResponse(data)


@_api_login_required
def api_analytics(request):
    deliverables = permitted_deliverables(request.user).filter(archived=False)
    by_status = list(deliverables.values('status').annotate(count=Count('id')).order_by('status'))
    by_unit = list(
        permitted_staff(request.user).filter(active=True)
        .values('unit__name').annotate(count=Count('id')).order_by('unit__name')
    )
    by_priority = list(permitted_tasks(request.user).values('priority').annotate(count=Count('id')).order_by('priority'))
    by_type = list(deliverables.values('type').annotate(count=Count('id')).order_by('type'))
    training = list(
        permitted_training(request.user).filter(archived=False)
        .values('validation_status').annotate(count=Count('id')).order_by('validation_status')
    )
    return JsonResponse({
        'ok': True, 'by_status': by_status, 'by_unit': by_unit,
        'by_priority': by_priority, 'by_type': by_type, 'training': training,
    })


@_api_login_required
@require_http_methods(['GET', 'POST'])
def api_records(request):
    if request.method == 'GET':
        qs = permitted_tasks(request.user).select_related('unit', 'assignee', 'category', 'created_by')
        q = (request.GET.get('search') or '').strip()
        cat = (request.GET.get('category') or '').strip()
        classification = (request.GET.get('classification') or '').strip()
        status = (request.GET.get('status') or '').strip()
        priority = (request.GET.get('priority') or '').strip()
        sort = request.GET.get('sort') or '-created_at'
        if q:
            qs = qs.filter(
                Q(title__icontains=q) | Q(description__icontains=q)
                | Q(assignee__username__icontains=q) | Q(unit__name__icontains=q)
            )
        if cat:
            if not cat.isdigit():
                return _json_error('Invalid category ID.', 400, {'category': ['Enter a valid unit ID.']})
            if not permitted_units(request.user).filter(pk=int(cat)).exists():
                return _json_error('Category is not in your scope.', 400, {'category': ['Out of scope.']})
            qs = qs.filter(unit_id=int(cat))
        if classification:
            if not classification.isdigit():
                return _json_error('Invalid classification ID.', 400)
            qs = qs.filter(category_id=int(classification))
        if status in TASK_STATUSES:
            qs = qs.filter(status=status)
        if priority in ('normal', 'high', 'critical'):
            qs = qs.filter(priority=priority)
        allowed_sorts = {
            'name': 'title', '-name': '-title', 'date': 'due_date', '-date': '-due_date',
            'due_date': 'due_date', '-due_date': '-due_date',
            'created_at': 'created_at', '-created_at': '-created_at',
            'status': 'status', '-status': '-status',
            'priority': 'priority', '-priority': '-priority',
        }
        qs = qs.order_by(allowed_sorts.get(sort, '-created_at'))
        try:
            requested_page = max(1, int(request.GET.get('page', '1')))
        except ValueError:
            requested_page = 1
        paginator = Paginator(qs, 10)
        page = paginator.get_page(requested_page)
        return JsonResponse({
            'ok': True,
            'results': [task_to_dict(t, request.user) for t in page.object_list],
            'page': page.number, 'pages': paginator.num_pages, 'total': paginator.count,
        })

    p = profile(request.user)
    if p.role not in (SUPERVISOR, MANAGER):
        return _json_error('Only a supervisor or unit manager can create records.', 403)
    try:
        payload = json.loads(request.body.decode('utf-8') or '{}')
    except (json.JSONDecodeError, UnicodeDecodeError):
        return _json_error('Invalid JSON payload.', 400)
    if not isinstance(payload, dict):
        return _json_error('JSON payload must be an object.', 400)
    if p.role == MANAGER:
        raw_unit = payload.get('unit')
        if raw_unit not in (None, ''):
            requested_unit = _int_or_none(raw_unit)
            if requested_unit is None:
                return _json_error('Invalid unit ID.', 400, {'unit': ['Enter a valid unit.']})
            if requested_unit != p.unit_id:
                return _json_error('Managers may only file records under their own unit.', 403)
    form = TaskForm(request.user, payload)
    if not form.is_valid():
        return _json_error('Validation failed', 400, _form_errors(form))
    with transaction.atomic():
        obj = form.save()
        log_audit(request.user, 'task_created', obj, detail=_task_payload_audit(obj), request=request)
        if obj.assignee_id:
            notify(
                obj.assignee, 'New task assigned',
                f'{obj.title} is due {obj.due_date}.',
                'task', reverse('tasks'),
                dedupe_key=f'task-{obj.pk}-assigned-{obj.assignee_id}',
            )
    return JsonResponse({'ok': True, 'record': task_to_dict(obj, request.user)}, status=201)


@_api_login_required
@require_http_methods(['GET', 'PUT', 'PATCH', 'DELETE'])
def api_record_detail(request, pk):
    obj = get_object_or_404(
        permitted_tasks(request.user).select_related('unit', 'assignee', 'category', 'created_by'), pk=pk
    )
    p = profile(request.user)
    if request.method == 'GET':
        history = list(
            AuditEvent.objects.filter(object_type='Task', object_id=obj.pk)
            .select_related('actor').order_by('-occurred_at')[:10]
            .values('event', 'occurred_at', 'actor__username', 'target_label')
        )
        for row in history:
            row['occurred_at'] = row['occurred_at'].isoformat()
        return JsonResponse({'ok': True, 'record': task_to_dict(obj, request.user), 'history': history})

    if request.method == 'DELETE':
        if p.role == SUPERVISOR:
            snapshot = _task_payload_audit(obj)
            with transaction.atomic():
                log_audit(
                    request.user, 'task_deleted', object_type='Task', object_id=obj.pk,
                    target_label=str(obj), detail=snapshot, request=request,
                )
                obj.delete()
            return JsonResponse({'ok': True, 'deleted': pk})
        if can_archive_task(request.user, obj):
            old = obj.status
            with transaction.atomic():
                obj.status = 'archived'
                obj.save()
                log_audit(
                    request.user, 'task_archived', obj,
                    previous={'status': old}, changes={'status': {'from': old, 'to': 'archived'}},
                    request=request,
                )
            return JsonResponse({'ok': True, 'archived': pk, 'record': task_to_dict(obj, request.user)})
        return _json_error('You cannot delete or archive this record.', 403)

    try:
        payload = json.loads(request.body.decode('utf-8') or '{}')
    except (json.JSONDecodeError, UnicodeDecodeError):
        return _json_error('Invalid JSON payload.', 400)
    if not isinstance(payload, dict):
        return _json_error('JSON payload must be an object.', 400)

    if p.role == STAFF:
        if obj.assignee_id != request.user.id:
            return _json_error('You may only update tasks assigned to you.', 403)
        if obj.status == 'archived':
            return _json_error('Archived tasks cannot be changed by staff.', 403)
        extras = set(payload) - {'status'}
        if extras:
            return _json_error('Staff may only update the task status.', 403)
        status = payload.get('status')
        if status == 'archived':
            return _json_error('Archiving requires manager permission.', 403)
        form = TaskStatusForm(obj, {'status': status})
        if not form.is_valid():
            return _json_error('Validation failed', 400, _form_errors(form))
        old = obj.status
        with transaction.atomic():
            obj.status = form.cleaned_data['status']
            obj.save()
            log_audit(
                request.user, 'task_status_' + obj.status, obj,
                previous={'status': old}, changes={'status': {'from': old, 'to': obj.status}},
                request=request,
            )
        return JsonResponse({'ok': True, 'record': task_to_dict(obj, request.user)})

    if p.role == MANAGER:
        raw_unit = payload.get('unit')
        if raw_unit not in (None, ''):
            requested_unit = _int_or_none(raw_unit)
            if requested_unit is None:
                return _json_error('Invalid unit ID.', 400, {'unit': ['Enter a valid unit.']})
            if requested_unit != p.unit_id:
                return _json_error('Managers may only edit records in their own unit.', 403)
    old = _task_payload_audit(obj)
    form = TaskForm(request.user, payload, instance=obj)
    if not form.is_valid():
        return _json_error('Validation failed', 400, _form_errors(form))
    if p.role == MANAGER and obj.unit_id and obj.unit_id != p.unit_id:
        return _json_error('Managers may only edit records in their own unit.', 403)
    with transaction.atomic():
        obj = form.save()
        new = _task_payload_audit(obj)
        log_audit(
            request.user, 'task_updated', obj,
            detail=new, previous=old,
            changes={k: {'from': old.get(k), 'to': new.get(k)} for k in new if old.get(k) != new.get(k)},
            request=request,
        )
        if obj.assignee_id and old.get('assignee') != obj.assignee_id:
            notify(obj.assignee, 'Task assigned', f'{obj.title} is due {obj.due_date}.', 'task', reverse('tasks'))
    return JsonResponse({'ok': True, 'record': task_to_dict(obj, request.user)})


@_api_login_required
@require_POST
def api_record_archive(request, pk):
    obj = get_object_or_404(permitted_tasks(request.user), pk=pk)
    if not can_archive_task(request.user, obj):
        return _json_error('You cannot archive this record.', 403)
    old = obj.status
    with transaction.atomic():
        obj.status = 'archived'
        obj.save()
        log_audit(request.user, 'task_archived', obj,
                  previous={'status': old}, changes={'status': {'from': old, 'to': 'archived'}}, request=request)
    return JsonResponse({'ok': True, 'record': task_to_dict(obj, request.user)})


@_api_login_required
@require_POST
def api_record_restore(request, pk):
    obj = get_object_or_404(permitted_tasks(request.user), pk=pk)
    if not can_archive_task(request.user, obj):
        return _json_error('You cannot restore this record.', 403)
    if obj.status != 'archived':
        return _json_error('Only archived records can be restored.', 400)
    with transaction.atomic():
        obj.status = 'pending'
        obj.save()
        log_audit(request.user, 'task_restored', obj,
                  previous={'status': 'archived'}, changes={'status': {'from': 'archived', 'to': 'pending'}},
                  request=request)
    return JsonResponse({'ok': True, 'record': task_to_dict(obj, request.user)})


@_api_login_required
@require_POST
def api_records_bulk(request):
    p = profile(request.user)
    try:
        payload = json.loads(request.body.decode('utf-8') or '{}')
    except (json.JSONDecodeError, UnicodeDecodeError):
        return _json_error('Invalid JSON payload.', 400)
    action = payload.get('action')
    ids = payload.get('ids')
    if action not in ('archive', 'restore', 'delete'):
        return _json_error('Invalid bulk action.', 400)
    if not isinstance(ids, list) or not ids:
        return _json_error('Select at least one record.', 400)
    clean_ids = []
    for value in ids:
        parsed = _int_or_none(value)
        if parsed is None:
            return _json_error('Invalid record ID.', 400)
        clean_ids.append(parsed)
    if action == 'delete' and p.role != SUPERVISOR:
        return _json_error('Permanent deletion is restricted to supervisors.', 403)
    qs = permitted_tasks(request.user).filter(pk__in=clean_ids)
    if qs.count() != len(set(clean_ids)):
        return _json_error('One or more records are outside your scope.', 403)
    objects = list(qs.select_related('unit'))
    if action in ('archive', 'restore'):
        for obj in objects:
            if not can_archive_task(request.user, obj):
                return _json_error('You cannot archive one or more selected records.', 403)
    with transaction.atomic():
        if action == 'delete':
            for obj in objects:
                log_audit(request.user, 'task_deleted', object_type='Task', object_id=obj.pk,
                          target_label=str(obj), detail=_task_payload_audit(obj), request=request)
                obj.delete()
        elif action == 'archive':
            for obj in objects:
                old = obj.status
                obj.status = 'archived'
                obj.save()
                log_audit(request.user, 'task_archived', obj, previous={'status': old},
                          changes={'status': {'from': old, 'to': 'archived'}}, request=request)
        else:
            for obj in objects:
                if obj.status != 'archived':
                    continue
                obj.status = 'pending'
                obj.save()
                log_audit(request.user, 'task_restored', obj, previous={'status': 'archived'},
                          changes={'status': {'from': 'archived', 'to': 'pending'}}, request=request)
    return JsonResponse({'ok': True, 'count': len(objects)})


@_api_login_required
def api_notifications(request):
    unread = Notification.objects.filter(user=request.user, is_read=False).count()
    qs = Notification.objects.filter(user=request.user).order_by('-created_at')[:20]
    return JsonResponse({
        'ok': True,
        'unread': unread,
        'results': [{
            'id': n.pk, 'title': n.title, 'message': n.message, 'category': n.category,
            'url': n.url or '', 'is_read': n.is_read,
            'created_at': n.created_at.isoformat(),
        } for n in qs],
    })


@_api_login_required
@require_POST
def api_notifications_read(request):
    try:
        payload = json.loads(request.body.decode('utf-8') or '{}')
    except (json.JSONDecodeError, UnicodeDecodeError):
        return _json_error('Invalid JSON payload.', 400)
    qs = Notification.objects.filter(user=request.user)
    if payload.get('all'):
        qs = qs.filter(is_read=False)
    else:
        notification_id = _int_or_none(payload.get('id'))
        if notification_id is None:
            return _json_error('Notification ID is required.', 400)
        qs = qs.filter(pk=notification_id)
    updated = qs.update(is_read=True, read_at=timezone.now())
    return JsonResponse({
        'ok': True, 'updated': updated,
        'unread': Notification.objects.filter(user=request.user, is_read=False).count(),
    })


@_api_login_required
def api_notifications_unread(request):
    return JsonResponse({
        'ok': True,
        'unread': Notification.objects.filter(user=request.user, is_read=False).count(),
    })


# ---------------------------------------------------------------------------
# Categories and units
# ---------------------------------------------------------------------------
@login_required
def categories_page(request):
    p = profile(request.user)
    units = permitted_units(request.user).annotate(
        staff_count=Count('staffrecord', distinct=True),
        task_count=Count('task', distinct=True),
        shift_count=Count('shiftcount', distinct=True),
    ).order_by('name')
    categories = (Category.objects.all() if can_manage_config(request.user) else Category.objects.filter(active=True))
    categories = categories.annotate(
        task_count=Count('tasks', distinct=True),
        training_count=Count('training', distinct=True),
    ).order_by('kind', 'name')
    type_counts = list(
        Deliverable.objects.filter(staff__in=permitted_staff(request.user))
        .values('type').annotate(count=Count('id')).order_by('type')
    )
    return render(request, 'core/categories.html', {
        'page_title': 'Categories & Units',
        'breadcrumb': 'Home / Categories',
        'units': units,
        'categories': categories,
        'type_counts': type_counts,
        'user_role': p.role,
        'can_write': can_manage_config(request.user),
    })


@login_required
@role_required(SUPERVISOR)
def unit_create(request):
    form = UnitForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        obj = form.save()
        log_audit(request.user, 'unit_created', obj, detail={'name': obj.name}, request=request)
        messages.success(request, f'Unit “{obj.name}” created.')
        return redirect('categories')
    return render(request, 'core/form.html', {
        'form': form, 'heading': 'Add unit', 'page_title': 'Add Unit',
        'breadcrumb': 'Home / Categories / Add Unit', 'cancel_url': reverse('categories'),
    })


@login_required
@role_required(SUPERVISOR)
def unit_edit(request, pk):
    obj = get_object_or_404(Unit, pk=pk)
    old = {'name': obj.name, 'description': obj.description, 'active': obj.active}
    form = UnitForm(request.POST or None, instance=obj)
    if request.method == 'POST' and form.is_valid():
        obj = form.save()
        log_audit(request.user, 'unit_updated', obj,
                  previous=old, changes={'name': obj.name, 'active': obj.active}, request=request)
        messages.success(request, 'Unit updated.')
        return redirect('categories')
    return render(request, 'core/form.html', {
        'form': form, 'heading': f'Edit unit — {obj.name}', 'page_title': 'Edit Unit',
        'breadcrumb': 'Home / Categories / Edit Unit', 'cancel_url': reverse('categories'),
    })


def _unit_references(unit):
    return {
        'staff': StaffRecord.objects.filter(unit=unit).count(),
        'active_staff': StaffRecord.objects.filter(unit=unit, active=True).count(),
        'profiles': Profile.objects.filter(unit=unit).count(),
        'tasks': Task.objects.filter(unit=unit).count(),
        'deliverables': Deliverable.objects.filter(staff__unit=unit).count(),
        'shifts': ShiftCount.objects.filter(unit=unit).count(),
    }


@login_required
@role_required(SUPERVISOR)
@require_POST
def unit_archive(request, pk):
    unit = get_object_or_404(Unit, pk=pk)
    refs = _unit_references(unit)
    if refs['active_staff'] or refs['tasks'] or refs['deliverables'] or refs['shifts'] or refs['profiles']:
        messages.error(request, 'This unit is referenced by active records. Reassign or archive those records first.')
        return redirect('categories')
    unit.active = False
    unit.save()
    log_audit(request.user, 'unit_archived', unit, detail=refs, request=request)
    messages.success(request, 'Unit deactivated.')
    return redirect('categories')


@login_required
@role_required(SUPERVISOR)
@require_POST
def unit_restore(request, pk):
    unit = get_object_or_404(Unit, pk=pk)
    unit.active = True
    unit.save()
    log_audit(request.user, 'unit_restored', unit, request=request)
    messages.success(request, 'Unit reactivated.')
    return redirect('categories')


@login_required
@role_required(SUPERVISOR)
@require_POST
def unit_delete(request, pk):
    unit = get_object_or_404(Unit, pk=pk)
    refs = _unit_references(unit)
    if any(refs.values()):
        messages.error(request, 'This unit is referenced by records and cannot be deleted. Deactivate it instead.')
        return redirect('categories')
    name = unit.name
    unit.delete()
    log_audit(request.user, 'unit_deleted', object_type='Unit', object_id=pk, target_label=name, request=request)
    messages.success(request, f'Unit “{name}” deleted.')
    return redirect('categories')


@login_required
@role_required(SUPERVISOR)
def category_create(request):
    form = CategoryForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        obj = form.save()
        log_audit(request.user, 'category_created', obj, detail={'name': obj.name, 'kind': obj.kind}, request=request)
        messages.success(request, 'Classification created.')
        return redirect('categories')
    return render(request, 'core/form.html', {
        'form': form, 'heading': 'Add classification', 'page_title': 'Add Classification',
        'breadcrumb': 'Home / Categories / Add Classification', 'cancel_url': reverse('categories'),
    })


@login_required
@role_required(SUPERVISOR)
def category_edit(request, pk):
    obj = get_object_or_404(Category, pk=pk)
    form = CategoryForm(request.POST or None, instance=obj)
    if request.method == 'POST' and form.is_valid():
        obj = form.save()
        log_audit(request.user, 'category_updated', obj, detail={'name': obj.name, 'kind': obj.kind}, request=request)
        messages.success(request, 'Classification updated.')
        return redirect('categories')
    return render(request, 'core/form.html', {
        'form': form, 'heading': f'Edit classification — {obj.name}', 'page_title': 'Edit Classification',
        'breadcrumb': 'Home / Categories / Edit Classification', 'cancel_url': reverse('categories'),
    })


@login_required
@role_required(SUPERVISOR)
@require_POST
def category_archive(request, pk):
    obj = get_object_or_404(Category, pk=pk)
    obj.active = False
    obj.save()
    log_audit(request.user, 'category_archived', obj, request=request)
    messages.success(request, 'Classification deactivated.')
    return redirect('categories')


@login_required
@role_required(SUPERVISOR)
@require_POST
def category_restore(request, pk):
    obj = get_object_or_404(Category, pk=pk)
    obj.active = True
    obj.save()
    log_audit(request.user, 'category_restored', obj, request=request)
    messages.success(request, 'Classification reactivated.')
    return redirect('categories')


@login_required
@role_required(SUPERVISOR)
@require_POST
def category_delete(request, pk):
    obj = get_object_or_404(Category, pk=pk)
    refs = Task.objects.filter(category=obj).count() + Training.objects.filter(category=obj).count()
    if refs:
        messages.error(request, 'This classification is referenced by records and cannot be deleted. Deactivate it instead.')
        return redirect('categories')
    name = obj.name
    obj.delete()
    log_audit(request.user, 'category_deleted', object_type='Category', object_id=pk, target_label=name, request=request)
    messages.success(request, 'Classification deleted.')
    return redirect('categories')


# ---------------------------------------------------------------------------
# Reports and exports
# ---------------------------------------------------------------------------
@login_required
def reports_page(request):
    data = request.GET.copy()
    if not data:
        data['kind'] = 'compliance'
    form = ReportFilterForm(request.user, data=data)
    report = None
    if form.is_valid():
        report = build_report(request.user, form.cleaned_data)
    return render(request, 'core/reports.html', {
        'page_title': 'Reports & Analytics',
        'breadcrumb': 'Home / Reports',
        'form': form,
        'report': report,
        'query_string': request.GET.urlencode(),
    })


@login_required
def report_export(request):
    data = request.GET.copy()
    fmt = (data.pop('format', ['csv'])[0] or 'csv').lower()
    if fmt not in ('csv', 'xlsx'):
        return HttpResponseBadRequest('Unsupported export format.')
    form = ReportFilterForm(request.user, data=data)
    if not form.is_valid():
        return HttpResponseBadRequest('Invalid report filters: ' + '; '.join(
            f'{field}: {", ".join(errors)}' for field, errors in form.errors.items()
        ))
    report = build_report(request.user, form.cleaned_data)
    log_audit(
        request.user, 'report_exported', object_type='Report', object_id=0,
        target_label=report['title'],
        detail={'kind': report['kind'], 'format': fmt, 'filters': {
            k: str(v) for k, v in form.cleaned_data.items() if k in ('start', 'end', 'unit', 'staff', 'status', 'type', 'horizon')
        }},
        request=request,
    )
    filename = report_filename(report['kind'], 'xlsx' if fmt == 'xlsx' else 'csv')
    if fmt == 'csv':
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        response.write('\ufeff')
        writer = csv.writer(response)
        writer.writerow([report['title']])
        writer.writerow(report['columns'])
        for row in report['rows']:
            writer.writerow([_safe_cell(value) for value in row])
        writer.writerow([])
        writer.writerow(['Summary'])
        for label, value in report['summary']:
            writer.writerow([label, _safe_cell(value)])
        return response

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report['kind'][:30]
    sheet.append([report['title']])
    sheet['A1'].font = Font(bold=True, size=14)
    sheet.append([])
    sheet.append(report['columns'])
    header_fill = PatternFill('solid', fgColor='DBEAFE')
    for cell in sheet[3]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in report['rows']:
        sheet.append([_safe_cell(value) for value in row])
    for column_cells in sheet.columns:
        length = max(len(str(cell.value or '')) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 45)
    summary_sheet = workbook.create_sheet('Summary')
    summary_sheet.append(['Measure', 'Value'])
    summary_sheet['A1'].font = Font(bold=True)
    summary_sheet['B1'].font = Font(bold=True)
    for label, value in report['summary']:
        summary_sheet.append([label, _safe_cell(value)])
    summary_sheet.column_dimensions['A'].width = 30
    summary_sheet.column_dimensions['B'].width = 20
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ---------------------------------------------------------------------------
# Users and settings
# ---------------------------------------------------------------------------
@login_required
@role_required(SUPERVISOR)
def users_page(request):
    qs = Profile.objects.select_related('user', 'unit').order_by('role', 'user__username')
    q = (request.GET.get('q') or '').strip()
    role_filter = request.GET.get('role') or ''
    unit_filter = request.GET.get('unit') or ''
    status_filter = request.GET.get('status') or ''
    if q:
        qs = qs.filter(
            Q(user__username__icontains=q) | Q(user__first_name__icontains=q)
            | Q(user__last_name__icontains=q) | Q(user__email__icontains=q)
        )
    if role_filter in (SUPERVISOR, MANAGER, STAFF):
        qs = qs.filter(role=role_filter)
    if unit_filter.isdigit():
        qs = qs.filter(unit_id=int(unit_filter))
    if status_filter == 'active':
        qs = qs.filter(user__is_active=True)
    elif status_filter == 'inactive':
        qs = qs.filter(user__is_active=False)
    page = _page(request, qs, 25)
    return render(request, 'core/users.html', {
        'page_title': 'User Management',
        'breadcrumb': 'Home / Users',
        'profiles': page,
        'page_obj': page,
        'units': Unit.objects.order_by('name'),
        'q': q, 'role_filter': role_filter, 'unit_filter': unit_filter, 'status_filter': status_filter,
    })


@login_required
@role_required(SUPERVISOR)
def user_create(request):
    form = UserCreateForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            user = form.save()
            log_audit(
                request.user, 'user_created', object_type='User', object_id=user.pk,
                target_label=user.username,
                detail={'username': user.username, 'role': form.cleaned_data['role'],
                        'unit': form.cleaned_data.get('unit').name if form.cleaned_data.get('unit') else None},
                request=request,
            )
        messages.success(request, f'User “{user.username}” created.')
        return redirect('users')
    return render(request, 'core/form.html', {
        'form': form, 'heading': 'Create user', 'page_title': 'Create User',
        'breadcrumb': 'Home / Users / Create', 'cancel_url': reverse('users'),
        'submit_label': 'Create user',
    })


@login_required
@role_required(SUPERVISOR)
def user_edit(request, pk):
    target = get_object_or_404(User.objects.select_related('nursing_profile'), pk=pk)
    old = {
        'username': target.username, 'first_name': target.first_name, 'last_name': target.last_name,
        'email': target.email, 'is_active': target.is_active,
    }
    form = UserUpdateForm(request.user, target, request.POST or None, instance=target)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            user = form.save()
            log_audit(
                request.user, 'user_updated', object_type='User', object_id=user.pk,
                target_label=user.username,
                detail={'role': form.cleaned_data['role'],
                        'unit': form.cleaned_data.get('unit').name if form.cleaned_data.get('unit') else None,
                        'is_active': user.is_active},
                previous=old,
                changes={k: {'from': old[k], 'to': getattr(user, k)} for k in old if old[k] != getattr(user, k)},
                request=request,
            )
        messages.success(request, 'User account updated.')
        return redirect('users')
    return render(request, 'core/form.html', {
        'form': form, 'heading': f'Edit user — {target.username}', 'page_title': 'Edit User',
        'breadcrumb': 'Home / Users / Edit', 'cancel_url': reverse('users'),
        'submit_label': 'Save account',
    })


@login_required
@role_required(SUPERVISOR)
@require_POST
def user_toggle(request, pk):
    target = get_object_or_404(User, pk=pk)
    if target.pk == request.user.pk:
        messages.error(request, 'You cannot deactivate your own account.')
        return redirect('users')
    target_profile = getattr(target, 'nursing_profile', None)
    if target.is_active and target_profile and target_profile.role == SUPERVISOR:
        if not Profile.objects.filter(role=SUPERVISOR, user__is_active=True).exclude(user=target).exists():
            messages.error(request, 'At least one active supervisor account must remain.')
            return redirect('users')
    target.is_active = not target.is_active
    target.save(update_fields=['is_active'])
    log_audit(request.user, 'user_activated' if target.is_active else 'user_deactivated',
              object_type='User', object_id=target.pk, target_label=target.username, request=request)
    messages.success(request, f'User “{target.username}” is now {"active" if target.is_active else "inactive"}.')
    return redirect('users')


@login_required
@role_required(SUPERVISOR)
@require_POST
def user_reset_password(request, pk):
    target = get_object_or_404(User, pk=pk)
    if not target.email:
        messages.error(request, 'This account has no email address. Add one before requesting a reset.')
        return redirect('users')
    form = PasswordResetForm({'email': target.email})
    if form.is_valid():
        try:
            form.save(
                request=request, use_https=request.is_secure(),
                from_email=settings.DEFAULT_FROM_EMAIL,
                email_template_name='registration/password_reset_email.html',
            )
            log_audit(request.user, 'password_reset_requested', object_type='User', object_id=target.pk,
                      target_label=target.username, request=request)
            messages.success(
                request,
                'A password reset message was generated using Django’s secure reset flow. '
                + ('In development it is printed to the server console.' if settings.DEBUG else ''),
            )
        except Exception:
            messages.error(request, 'The reset message could not be sent. Check the email configuration.')
    else:
        messages.error(request, 'A valid email address is required for password reset.')
    return redirect('users')


@login_required
def settings_page(request):
    p = profile(request.user)
    form = ProfileSelfForm(instance=request.user)
    if request.method == 'POST':
        old = {'first_name': request.user.first_name, 'last_name': request.user.last_name, 'email': request.user.email}
        form = ProfileSelfForm(request.POST, instance=request.user)
        if form.is_valid():
            with transaction.atomic():
                form.save()
                log_audit(request.user, 'profile_updated', object_type='User', object_id=request.user.pk,
                          target_label=request.user.username,
                          previous=old,
                          changes={k: {'from': old[k], 'to': getattr(request.user, k)}
                                   for k in old if old[k] != getattr(request.user, k)},
                          request=request)
            messages.success(request, 'Profile updated.')
            return redirect('settings')
    my_staff = getattr(request.user, 'staff_record', None)
    recent = AuditEvent.objects.filter(actor=request.user).order_by('-occurred_at')[:8]
    return render(request, 'core/settings.html', {
        'page_title': 'Settings',
        'breadcrumb': 'Home / Settings',
        'my_profile': p,
        'form': form,
        'my_staff': my_staff,
        'recent_activity': recent,
    })


# ---------------------------------------------------------------------------
# Staff 201
# ---------------------------------------------------------------------------
@login_required
def staff_directory(request):
    p = profile(request.user)
    qs = permitted_staff(request.user).select_related('user', 'unit')
    q = (request.GET.get('q') or '').strip()
    unit_filter = (request.GET.get('unit') or '').strip()
    status_filter = request.GET.get('status') or ''
    expiry_filter = request.GET.get('expiring') or ''
    if q:
        qs = qs.filter(
            Q(user__username__icontains=q) | Q(user__first_name__icontains=q)
            | Q(user__last_name__icontains=q) | Q(employee_number__icontains=q)
            | Q(position__icontains=q) | Q(prc_number__icontains=q)
        )
    if unit_filter.isdigit() and permitted_units(request.user).filter(pk=int(unit_filter)).exists():
        qs = qs.filter(unit_id=int(unit_filter))
    if status_filter == 'active':
        qs = qs.filter(active=True)
    elif status_filter == 'inactive':
        qs = qs.filter(active=False)
    if expiry_filter == 'soon':
        today = timezone.localdate()
        qs = qs.filter(active=True, prc_expiry__isnull=False,
                       prc_expiry__gte=today, prc_expiry__lte=today + timedelta(days=90))
    qs = qs.order_by('unit__name', 'user__last_name', 'user__first_name')
    page = _page(request, qs, 20)
    rows = [(s, can_view_sensitive_staff(request.user, s)) for s in page.object_list]
    return render(request, 'core/staff.html', {
        'staff': rows,
        'page_obj': page,
        'user_role': p.role,
        'can_add': can_create_staff(request.user),
        'units': permitted_units(request.user).order_by('name'),
        'q': q, 'unit_filter': unit_filter, 'status_filter': status_filter, 'expiry_filter': expiry_filter,
        'page_title': '201 Files',
        'breadcrumb': 'Home / 201 Files',
    })


@login_required
@role_required(SUPERVISOR)
def staff_create(request):
    form = StaffRecordForm(request.user, request.POST or None)
    if request.method == 'POST' and form.is_valid():
        obj = form.save()
        log_audit(request.user, 'staff_created', obj,
                  detail={'employee_number': obj.employee_number, 'unit': obj.unit.name,
                          'position': obj.position, 'employment_status': obj.employment_status},
                  request=request)
        messages.success(request, 'Staff 201 file created.')
        return redirect('staff_detail', pk=obj.pk)
    return render(request, 'core/form.html', {
        'form': form, 'heading': 'Create staff 201 file', 'page_title': 'Create 201 File',
        'breadcrumb': 'Home / 201 Files / Create', 'cancel_url': reverse('staff'),
        'submit_label': 'Create 201 file',
    })


@login_required
def staff_detail(request, pk):
    staff = get_object_or_404(permitted_staff(request.user).select_related('user', 'unit'), pk=pk)
    p = profile(request.user)
    show_sensitive = can_view_sensitive_staff(request.user, staff)
    today = timezone.localdate()
    license_days = (staff.prc_expiry - today).days if staff.prc_expiry else None
    audit_history = AuditEvent.objects.filter(object_type='StaffRecord', object_id=staff.pk).select_related('actor')[:15]
    access_history = AuditEvent.objects.filter(actor=staff.user).select_related('actor').order_by('-occurred_at')[:10]
    return render(request, 'core/staff_detail.html', {
        's': staff,
        'user_role': p.role,
        'show_sensitive': show_sensitive,
        'can_edit': can_edit_staff(request.user, staff),
        'can_toggle': can_toggle_staff(request.user, staff),
        'license_days': license_days,
        'trainings': staff.training.select_related('category').order_by('-date')[:20],
        'deliverables': staff.deliverables.order_by('-due_date')[:20],
        'audit_history': audit_history,
        'access_history': access_history,
        'page_title': '201 File Detail',
        'breadcrumb': 'Home / 201 Files / Detail',
    })


@login_required
def staff_edit(request, pk):
    staff = get_object_or_404(permitted_staff(request.user).select_related('user', 'unit'), pk=pk)
    if not can_edit_staff(request.user, staff):
        raise PermissionDenied('You cannot edit this 201 file.')
    old = {
        'employee_number': staff.employee_number, 'unit': staff.unit_id, 'position': staff.position,
        'employment_status': staff.employment_status, 'active': staff.active,
    }
    form = StaffRecordForm(request.user, request.POST or None, instance=staff)
    if request.method == 'POST' and form.is_valid():
        obj = form.save()
        new = {
            'employee_number': obj.employee_number, 'unit': obj.unit_id, 'position': obj.position,
            'employment_status': obj.employment_status, 'active': obj.active,
        }
        log_audit(request.user, 'staff_updated', obj,
                  detail=new, previous=old,
                  changes={k: {'from': old[k], 'to': new[k]} for k in old if old[k] != new[k]},
                  request=request)
        messages.success(request, '201 file updated.')
        return redirect('staff_detail', pk=obj.pk)
    return render(request, 'core/form.html', {
        'form': form, 'heading': f'Edit 201 file — {staff}', 'page_title': 'Edit 201 File',
        'breadcrumb': 'Home / 201 Files / Edit', 'cancel_url': reverse('staff_detail', args=[staff.pk]),
        'submit_label': 'Save 201 file',
    })


@login_required
@require_POST
def staff_toggle(request, pk):
    staff = get_object_or_404(permitted_staff(request.user), pk=pk)
    if not can_toggle_staff(request.user, staff):
        raise PermissionDenied('You cannot change this staff record.')
    staff.active = not staff.active
    staff.save()
    log_audit(request.user, 'staff_activated' if staff.active else 'staff_deactivated', staff,
              previous={'active': not staff.active}, changes={'active': {'from': not staff.active, 'to': staff.active}},
              request=request)
    messages.success(request, f'{staff} is now {"active" if staff.active else "inactive"}.')
    return redirect('staff_detail', pk=staff.pk)


# ---------------------------------------------------------------------------
# Deliverables
# ---------------------------------------------------------------------------
@login_required
def deliverables(request):
    p = profile(request.user)
    qs = permitted_deliverables(request.user).select_related('staff__user', 'staff__unit', 'reviewer')
    f_type = request.GET.get('type', '')
    f_status = request.GET.get('status', '')
    f_overdue = request.GET.get('overdue', '')
    f_unit = request.GET.get('unit', '')
    f_staff = request.GET.get('staff', '')
    f_period = (request.GET.get('period') or '').strip()
    show_archived = request.GET.get('archived') == '1'
    today = timezone.localdate()
    qs = qs.filter(archived=show_archived)
    if f_type in ('DTR', 'DAR', 'IPCR', 'OTHER'):
        qs = qs.filter(type=f_type)
    if f_status in DELIVERABLE_STATUSES:
        qs = qs.filter(status=f_status)
    if f_overdue == '1':
        qs = qs.filter(due_date__lt=today).exclude(status__in=['verified', 'approved'])
    if f_unit.isdigit() and permitted_units(request.user).filter(pk=int(f_unit)).exists():
        qs = qs.filter(staff__unit_id=int(f_unit))
    if f_staff.isdigit() and permitted_staff(request.user).filter(pk=int(f_staff)).exists():
        qs = qs.filter(staff_id=int(f_staff))
    if f_period:
        qs = qs.filter(period__icontains=f_period)
    qs = qs.order_by('due_date', 'type')
    page = _page(request, qs, 25)
    rows = [(d, d.is_overdue) for d in page.object_list]
    return render(request, 'core/deliverables.html', {
        'rows': rows,
        'page_obj': page,
        'f_type': f_type, 'f_status': f_status, 'f_overdue': f_overdue,
        'f_unit': f_unit, 'f_staff': f_staff, 'f_period': f_period,
        'show_archived': show_archived,
        'units': permitted_units(request.user).order_by('name'),
        'staff_options': permitted_staff(request.user).select_related('user').order_by('user__username'),
        'user_role': p.role,
        'can_create': p.role in (SUPERVISOR, MANAGER),
        'today': today,
        'page_title': 'Deliverables',
        'breadcrumb': 'Home / Deliverables',
    })


@login_required
@role_required(SUPERVISOR, MANAGER)
def deliverable_create(request):
    form = DeliverableForm(request.user, request.POST or None)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            obj = form.save()
            log_audit(request.user, 'deliverable_created', obj,
                      detail={'type': obj.type, 'period': obj.period, 'due_date': obj.due_date.isoformat(),
                              'staff': obj.staff_id}, request=request)
            notify(obj.staff.user, 'New deliverable requirement',
                   f'{obj.type} for {obj.period} is due {obj.due_date}.',
                   'deadline', reverse('deliverables'),
                   dedupe_key=f'deliv-{obj.pk}-created')
        messages.success(request, 'Deliverable requirement created.')
        return redirect('deliverables')
    return render(request, 'core/form.html', {
        'form': form, 'heading': 'Create deliverable', 'page_title': 'Create Deliverable',
        'breadcrumb': 'Home / Deliverables / Create', 'cancel_url': reverse('deliverables'),
        'submit_label': 'Create deliverable',
    })


@login_required
@role_required(SUPERVISOR, MANAGER)
def deliverable_bulk_create(request):
    form = BulkDeliverableForm(request.user, request.POST or None)
    result = None
    if request.method == 'POST' and form.is_valid():
        staff_list = list(form.cleaned_data['staff'])
        periods = form.cleaned_data['periods']
        dtype = form.cleaned_data['type']
        created = 0
        skipped = 0
        with transaction.atomic():
            for staff in staff_list:
                for period, due_date in periods:
                    obj, was_created = Deliverable.objects.get_or_create(
                        staff=staff, type=dtype, period=period,
                        defaults={'due_date': due_date, 'status': 'pending', 'created_by': request.user},
                    )
                    if not was_created:
                        skipped += 1
                        continue
                    created += 1
                    log_audit(request.user, 'deliverable_bulk_created', obj,
                              detail={'period': period, 'type': dtype, 'staff': staff.pk}, request=request)
                    notify(staff.user, 'New deliverable requirement',
                           f'{dtype} for {period} is due {due_date}.', 'deadline', reverse('deliverables'),
                           dedupe_key=f'deliv-{obj.pk}-created')
        result = {'created': created, 'skipped': skipped, 'staff': len(staff_list), 'periods': len(periods)}
        messages.success(request, f'{created} deliverable(s) created; {skipped} duplicate(s) skipped.')
    return render(request, 'core/deliverable_bulk_form.html', {
        'form': form, 'result': result, 'page_title': 'Bulk Create Deliverables',
        'breadcrumb': 'Home / Deliverables / Bulk Create', 'cancel_url': reverse('deliverables'),
    })


@login_required
def deliverable_detail(request, pk):
    d = get_object_or_404(permitted_deliverables(request.user).select_related('staff__user', 'staff__unit'), pk=pk)
    history = AuditEvent.objects.filter(object_type='Deliverable', object_id=d.pk).select_related('actor').order_by('occurred_at')
    return render(request, 'core/deliverable_detail.html', {
        'd': d,
        'history': history,
        'can_submit': can_submit_deliverable(request.user, d),
        'can_review': can_review_deliverable(request.user, d) and d.status in ('submitted', 'review', 'verified'),
        'can_edit': can_manage_deliverable(request.user, d) and d.status != 'approved',
        'can_archive': can_manage_deliverable(request.user, d),
        'can_delete': role(request.user) == SUPERVISOR,
        'page_title': 'Deliverable Detail',
        'breadcrumb': 'Home / Deliverables / Detail',
    })


@login_required
@role_required(SUPERVISOR, MANAGER)
def deliverable_edit(request, pk):
    d = get_object_or_404(permitted_deliverables(request.user).select_related('staff__user', 'staff__unit'), pk=pk)
    if not can_manage_deliverable(request.user, d):
        raise PermissionDenied('You cannot edit this deliverable.')
    old = {'type': d.type, 'period': d.period, 'due_date': d.due_date.isoformat(), 'staff': d.staff_id}
    form = DeliverableForm(request.user, request.POST or None, instance=d)
    if request.method == 'POST' and form.is_valid():
        obj = form.save()
        log_audit(request.user, 'deliverable_updated', obj,
                  detail={'type': obj.type, 'period': obj.period, 'due_date': obj.due_date.isoformat()},
                  previous=old, request=request)
        messages.success(request, 'Deliverable updated.')
        return redirect('deliverable_detail', pk=obj.pk)
    return render(request, 'core/form.html', {
        'form': form, 'heading': f'Edit deliverable — {d.type} {d.period}',
        'page_title': 'Edit Deliverable', 'breadcrumb': 'Home / Deliverables / Edit',
        'cancel_url': reverse('deliverable_detail', args=[d.pk]), 'submit_label': 'Save deliverable',
    })


@login_required
def submit(request, pk):
    d = get_object_or_404(permitted_deliverables(request.user).select_related('staff__user'), pk=pk)
    if not can_submit_deliverable(request.user, d):
        raise PermissionDenied('This deliverable is not open for your submission.')
    form = SubmissionForm(request.POST or None, instance=d)
    if request.method == 'POST' and form.is_valid():
        try:
            submit_deliverable(d, request.user, form.cleaned_data['document_reference'], request=request)
        except ValidationError as exc:
            form.add_error(None, exc)
        else:
            messages.success(request, 'Deliverable submitted for verification.')
            return redirect('deliverables')
    return render(request, 'core/form.html', {
        'form': form, 'heading': f'Submit {d.type} — {d.period}',
        'page_title': 'Submit Deliverable', 'breadcrumb': 'Home / Deliverables / Submit',
        'cancel_url': reverse('deliverable_detail', args=[d.pk]), 'submit_label': 'Submit for verification',
        'object': d,
    })


@login_required
def review(request, pk):
    d = get_object_or_404(permitted_deliverables(request.user).select_related('staff__user', 'staff__unit'), pk=pk)
    if not can_review_deliverable(request.user, d):
        raise PermissionDenied('You cannot review this deliverable.')
    form = ReviewForm(d, request.user, request.POST or None)
    if request.method == 'POST' and form.is_valid():
        try:
            transition_deliverable(d, request.user, form.cleaned_data['decision'],
                                   form.cleaned_data.get('remarks', ''), request=request)
        except ValidationError as exc:
            form.add_error(None, exc)
        else:
            messages.success(request, 'Workflow action recorded.')
            return redirect('deliverable_detail', pk=d.pk)
    history = AuditEvent.objects.filter(object_type='Deliverable', object_id=d.pk).select_related('actor').order_by('occurred_at')
    return render(request, 'core/review.html', {
        'form': form, 'd': d, 'history': history,
        'page_title': 'Review Deliverable', 'breadcrumb': 'Home / Deliverables / Review',
        'cancel_url': reverse('deliverable_detail', args=[d.pk]),
    })


@login_required
@require_POST
def deliverable_archive(request, pk):
    d = get_object_or_404(permitted_deliverables(request.user), pk=pk)
    if not can_manage_deliverable(request.user, d):
        raise PermissionDenied('You cannot archive this deliverable.')
    d.archived = True
    d.save()
    log_audit(request.user, 'deliverable_archived', d, request=request)
    messages.success(request, 'Deliverable archived.')
    return redirect('deliverables')


@login_required
@require_POST
def deliverable_restore(request, pk):
    d = get_object_or_404(permitted_deliverables(request.user), pk=pk)
    if not can_manage_deliverable(request.user, d):
        raise PermissionDenied('You cannot restore this deliverable.')
    d.archived = False
    d.save()
    log_audit(request.user, 'deliverable_restored', d, request=request)
    messages.success(request, 'Deliverable restored.')
    return redirect('deliverable_detail', pk=d.pk)


@login_required
@role_required(SUPERVISOR)
@require_POST
def deliverable_delete(request, pk):
    d = get_object_or_404(permitted_deliverables(request.user), pk=pk)
    snapshot = {'type': d.type, 'period': d.period, 'staff': d.staff_id, 'status': d.status}
    log_audit(request.user, 'deliverable_deleted', object_type='Deliverable', object_id=d.pk,
              target_label=str(d), detail=snapshot, request=request)
    d.delete()
    messages.success(request, 'Deliverable permanently deleted.')
    return redirect('deliverables')


# ---------------------------------------------------------------------------
# Training / LDNI
# ---------------------------------------------------------------------------
@login_required
def training(request):
    p = profile(request.user)
    qs = permitted_training(request.user).select_related('staff__user', 'staff__unit', 'category')
    q = (request.GET.get('q') or '').strip()
    f_unit = request.GET.get('unit') or ''
    f_status = request.GET.get('status') or ''
    f_validation = request.GET.get('validation') or ''
    show_archived = request.GET.get('archived') == '1'
    qs = qs.filter(archived=show_archived)
    if q:
        qs = qs.filter(Q(title__icontains=q) | Q(provider__icontains=q)
                       | Q(staff__user__username__icontains=q))
    if f_unit.isdigit() and permitted_units(request.user).filter(pk=int(f_unit)).exists():
        qs = qs.filter(staff__unit_id=int(f_unit))
    if f_status in TRAINING_STATUSES:
        qs = qs.filter(status=f_status)
    if f_validation in VALIDATION_STATUSES:
        qs = qs.filter(validation_status=f_validation)
    qs = qs.order_by('-date')
    page = _page(request, qs, 25)
    summary = qs.aggregate(hours=Sum('hours'))
    return render(request, 'core/training.html', {
        'training': page,
        'page_obj': page,
        'total_hours': summary['hours'] or 0,
        'validated_count': qs.filter(validation_status='validated').count(),
        'pending_count': qs.filter(validation_status='pending').count(),
        'units': permitted_units(request.user).order_by('name'),
        'q': q, 'f_unit': f_unit, 'f_status': f_status, 'f_validation': f_validation,
        'show_archived': show_archived,
        'user_role': p.role,
        'can_add': p.role in (SUPERVISOR, MANAGER) or bool(getattr(request.user, 'staff_record', None)),
        'page_title': 'LDNI / Training',
        'breadcrumb': 'Home / Training',
    })


@login_required
def add_training(request):
    p = profile(request.user)
    if p.role == STAFF and not getattr(request.user, 'staff_record', None):
        messages.error(request, 'Your account has no 201 file. Ask a supervisor to create one first.')
        return redirect('training')
    form = TrainingForm(request.user, request.POST or None)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            obj = form.save()
            log_audit(request.user, 'training_created', obj,
                      detail={'title': obj.title, 'date': obj.date.isoformat(), 'hours': str(obj.hours),
                              'validation_status': obj.validation_status}, request=request)
            if obj.validation_status == 'pending':
                notify_roles([MANAGER], 'Training awaiting validation',
                             f'{obj.staff} · {obj.title}', 'training', reverse('training'),
                             unit=obj.staff.unit_id, dedupe_suffix=f'training-{obj.pk}-pending')
        messages.success(request, 'Training record saved.')
        return redirect('training')
    return render(request, 'core/form.html', {
        'form': form, 'heading': 'Record training', 'page_title': 'Record Training',
        'breadcrumb': 'Home / Training / Add', 'cancel_url': reverse('training'),
        'submit_label': 'Save training',
    })


@login_required
def training_detail(request, pk):
    t = get_object_or_404(permitted_training(request.user).select_related('staff__user', 'staff__unit'), pk=pk)
    history = AuditEvent.objects.filter(object_type='Training', object_id=t.pk).select_related('actor').order_by('occurred_at')
    return render(request, 'core/training_detail.html', {
        't': t, 'history': history,
        'can_edit': can_edit_training(request.user, t),
        'can_validate': can_validate_training(request.user, t),
        'can_archive': can_archive_training(request.user, t),
        'can_delete': can_delete_training(request.user, t),
        'page_title': 'Training Detail', 'breadcrumb': 'Home / Training / Detail',
    })


@login_required
def training_edit(request, pk):
    t = get_object_or_404(permitted_training(request.user).select_related('staff__user'), pk=pk)
    if not can_edit_training(request.user, t):
        raise PermissionDenied('You cannot edit this training record.')
    old = {'title': t.title, 'date': t.date.isoformat(), 'hours': str(t.hours), 'provider': t.provider,
           'status': t.status, 'validation_status': t.validation_status}
    form = TrainingForm(request.user, request.POST or None, instance=t)
    if request.method == 'POST' and form.is_valid():
        obj = form.save()
        log_audit(request.user, 'training_updated', obj, detail={
            'title': obj.title, 'date': obj.date.isoformat(), 'hours': str(obj.hours),
            'status': obj.status, 'validation_status': obj.validation_status,
        }, previous=old, request=request)
        messages.success(request, 'Training record updated.')
        return redirect('training_detail', pk=obj.pk)
    return render(request, 'core/form.html', {
        'form': form, 'heading': f'Edit training — {t.title}', 'page_title': 'Edit Training',
        'breadcrumb': 'Home / Training / Edit', 'cancel_url': reverse('training_detail', args=[t.pk]),
        'submit_label': 'Save training',
    })


@login_required
def training_validate(request, pk):
    t = get_object_or_404(permitted_training(request.user).select_related('staff__user'), pk=pk)
    if not can_validate_training(request.user, t):
        raise PermissionDenied('You cannot validate this training record.')
    form = TrainingValidationForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        try:
            validate_training_service(t, request.user, form.cleaned_data['decision'],
                                      form.cleaned_data.get('remarks', ''), request=request)
        except ValidationError as exc:
            form.add_error(None, exc)
        else:
            messages.success(request, 'Training validation recorded.')
            return redirect('training_detail', pk=t.pk)
    return render(request, 'core/form.html', {
        'form': form, 'heading': f'Validate training — {t.title}', 'page_title': 'Validate Training',
        'breadcrumb': 'Home / Training / Validate', 'cancel_url': reverse('training_detail', args=[t.pk]),
        'submit_label': 'Save decision', 'object': t,
    })


@login_required
@require_POST
def training_archive(request, pk):
    t = get_object_or_404(permitted_training(request.user), pk=pk)
    if not can_archive_training(request.user, t):
        raise PermissionDenied('You cannot archive this training record.')
    t.archived = True
    t.save()
    log_audit(request.user, 'training_archived', t, request=request)
    messages.success(request, 'Training record archived.')
    return redirect('training')


@login_required
@require_POST
def training_restore(request, pk):
    t = get_object_or_404(permitted_training(request.user), pk=pk)
    if not can_archive_training(request.user, t):
        raise PermissionDenied('You cannot restore this training record.')
    t.archived = False
    t.save()
    log_audit(request.user, 'training_restored', t, request=request)
    messages.success(request, 'Training record restored.')
    return redirect('training_detail', pk=t.pk)


@login_required
@role_required(SUPERVISOR)
@require_POST
def training_delete(request, pk):
    t = get_object_or_404(permitted_training(request.user), pk=pk)
    log_audit(request.user, 'training_deleted', object_type='Training', object_id=t.pk,
              target_label=str(t), detail={'title': t.title, 'staff': t.staff_id}, request=request)
    t.delete()
    messages.success(request, 'Training record permanently deleted.')
    return redirect('training')


# ---------------------------------------------------------------------------
# Tasks / deadlines
# ---------------------------------------------------------------------------
@login_required
def tasks_view(request):
    p = profile(request.user)
    qs = permitted_tasks(request.user).select_related('unit', 'assignee', 'category')
    f = request.GET.get('filter', 'all')
    today = timezone.localdate()
    soon = today + timedelta(days=7)
    if f == 'overdue':
        qs = qs.filter(due_date__lt=today).exclude(status__in=['completed', 'archived'])
    elif f == 'soon':
        qs = qs.filter(due_date__gte=today, due_date__lte=soon).exclude(status__in=['completed', 'archived'])
    elif f == 'open':
        qs = qs.exclude(status__in=['completed', 'archived'])
    rows = []
    for t in qs.order_by('due_date', '-priority')[:200]:
        flag = ('done' if t.status in ('completed', 'archived')
                else 'overdue' if t.due_date < today
                else 'soon' if t.due_date <= soon
                else 'ok')
        rows.append((t, flag))
    return render(request, 'core/tasks.html', {
        'rows': rows, 'f': f, 'today': today,
        'can_add': p.role in (SUPERVISOR, MANAGER),
        'user_role': p.role,
        'page_title': 'Deadlines',
        'breadcrumb': 'Home / Deadlines',
    })


@login_required
@role_required(SUPERVISOR, MANAGER)
def add_task(request):
    form = TaskForm(request.user, request.POST or None)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            obj = form.save()
            log_audit(request.user, 'task_created', obj, detail=_task_payload_audit(obj), request=request)
            if obj.assignee_id:
                notify(obj.assignee, 'New task assigned', f'{obj.title} is due {obj.due_date}.',
                       'task', reverse('tasks'), dedupe_key=f'task-{obj.pk}-assigned-{obj.assignee_id}')
        messages.success(request, 'Task created.')
        return redirect('tasks')
    return render(request, 'core/form.html', {
        'form': form, 'heading': 'Assign task / deadline', 'page_title': 'Assign Task',
        'breadcrumb': 'Home / Deadlines / Add', 'cancel_url': reverse('tasks'), 'submit_label': 'Create task',
    })


@login_required
def task_edit(request, pk):
    t = get_object_or_404(permitted_tasks(request.user).select_related('unit'), pk=pk)
    if not can_edit_task(request.user, t) or role(request.user) == STAFF:
        raise PermissionDenied('Use the status control to update tasks assigned to you.')
    old = _task_payload_audit(t)
    form = TaskForm(request.user, request.POST or None, instance=t)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            obj = form.save()
            new = _task_payload_audit(obj)
            log_audit(request.user, 'task_updated', obj, detail=new, previous=old,
                      changes={k: {'from': old.get(k), 'to': new.get(k)} for k in new if old.get(k) != new.get(k)},
                      request=request)
        messages.success(request, 'Task updated.')
        return redirect('tasks')
    return render(request, 'core/form.html', {
        'form': form, 'heading': f'Edit task — {t.title}', 'page_title': 'Edit Task',
        'breadcrumb': 'Home / Deadlines / Edit', 'cancel_url': reverse('tasks'), 'submit_label': 'Save task',
    })


@login_required
@require_POST
def task_status(request, pk):
    t = get_object_or_404(permitted_tasks(request.user).select_related('unit'), pk=pk)
    if not can_edit_task(request.user, t):
        raise PermissionDenied('You cannot update this task.')
    p = profile(request.user)
    requested = request.POST.get('status', '')
    if p.role == STAFF and t.status == 'archived':
        raise PermissionDenied('Archived tasks cannot be changed by staff.')
    if p.role == STAFF:
        form = TaskStatusForm(t, {'status': requested})
    else:
        form = TaskStatusForm(t, {'status': requested})
        if requested == 'archived':
            if not can_archive_task(request.user, t):
                raise PermissionDenied('You cannot archive this task.')
            t.status = 'archived'
            t.save()
            log_audit(request.user, 'task_archived', t, request=request)
            messages.success(request, 'Task archived.')
            return redirect('tasks')
    if form.is_valid():
        old = t.status
        t.status = form.cleaned_data['status']
        t.save()
        log_audit(request.user, 'task_status_' + t.status, t,
                  previous={'status': old}, changes={'status': {'from': old, 'to': t.status}}, request=request)
        messages.success(request, 'Task status updated.')
    else:
        messages.error(request, 'Invalid task status.')
    return redirect('tasks')


# ---------------------------------------------------------------------------
# Shifts
# ---------------------------------------------------------------------------
@login_required
@role_required(SUPERVISOR, MANAGER)
def shifts(request):
    p = profile(request.user)
    qs = permitted_shifts(request.user).select_related('unit', 'entered_by')
    f_unit = request.GET.get('unit') or ''
    f_shift = request.GET.get('shift') or ''
    start = _date_or_none(request.GET.get('start'))
    end = _date_or_none(request.GET.get('end'))
    month_value = request.GET.get('month') or ''
    year, month = _parse_month(month_value)
    first, last = _month_bounds(year, month)
    if not start:
        start = first
    if not end:
        end = last
    if f_unit.isdigit() and permitted_units(request.user).filter(pk=int(f_unit)).exists():
        qs = qs.filter(unit_id=int(f_unit))
    if f_shift in ('day', 'night'):
        qs = qs.filter(shift=f_shift)
    qs = qs.filter(date__gte=start, date__lte=end)
    summary = {
        'records': qs.count(),
        'day': qs.filter(shift='day').aggregate(total=Sum('on_duty'))['total'] or 0,
        'night': qs.filter(shift='night').aggregate(total=Sum('on_duty'))['total'] or 0,
    }
    page = _page(request, qs.order_by('-date', 'unit__name', 'shift'), 31)
    calendar_rows = []
    if month_value:
        entries = {}
        for shift in qs:
            entries.setdefault(shift.date, []).append(shift)
        cal = calendar.Calendar(firstweekday=0)
        for week in cal.monthdatescalendar(year, month):
            cells = []
            for day in week:
                cells.append({
                    'date': day,
                    'in_month': day.month == month,
                    'entries': entries.get(day, []),
                })
            calendar_rows.append(cells)
    return render(request, 'core/shifts.html', {
        'shifts': page,
        'page_obj': page,
        'summary': summary,
        'units': permitted_units(request.user).order_by('name'),
        'f_unit': f_unit, 'f_shift': f_shift,
        'start': start, 'end': end,
        'month_value': f'{year:04d}-{month:02d}' if month_value else '',
        'month_label': date(year, month, 1).strftime('%B %Y'),
        'prev_month': (date(year, month, 1) - timedelta(days=1)).strftime('%Y-%m'),
        'next_month': (date(year, month, 28) + timedelta(days=4)).strftime('%Y-%m'),
        'calendar_rows': calendar_rows,
        'user_role': p.role,
        'page_title': 'Shift Counts',
        'breadcrumb': 'Home / Shifts',
    })


@login_required
@role_required(SUPERVISOR, MANAGER)
def add_shift(request):
    form = ShiftForm(request.user, request.POST or None)
    if request.method == 'POST' and form.is_valid():
        obj = form.save(commit=False)
        obj.entered_by = request.user
        obj.save()
        log_audit(request.user, 'shift_created', obj,
                  detail={'unit': obj.unit.name, 'date': obj.date.isoformat(), 'shift': obj.shift,
                          'on_duty': obj.on_duty}, request=request)
        messages.success(request, 'Shift count saved.')
        return redirect('shifts')
    return render(request, 'core/form.html', {
        'form': form, 'heading': 'Enter shift count', 'page_title': 'Add Shift',
        'breadcrumb': 'Home / Shifts / Add', 'cancel_url': reverse('shifts'), 'submit_label': 'Save shift count',
    })


@login_required
@role_required(SUPERVISOR, MANAGER)
def shift_edit(request, pk):
    shift = get_object_or_404(permitted_shifts(request.user).select_related('unit'), pk=pk)
    old = {'unit': shift.unit_id, 'date': shift.date.isoformat(), 'shift': shift.shift, 'on_duty': shift.on_duty}
    form = ShiftForm(request.user, request.POST or None, instance=shift)
    if request.method == 'POST' and form.is_valid():
        obj = form.save()
        log_audit(request.user, 'shift_updated', obj,
                  detail={'unit': obj.unit.name, 'date': obj.date.isoformat(), 'shift': obj.shift, 'on_duty': obj.on_duty},
                  previous=old, request=request)
        messages.success(request, 'Shift count updated.')
        return redirect('shifts')
    return render(request, 'core/form.html', {
        'form': form, 'heading': f'Edit shift — {shift.unit} {shift.date}',
        'page_title': 'Edit Shift', 'breadcrumb': 'Home / Shifts / Edit',
        'cancel_url': reverse('shifts'), 'submit_label': 'Save shift count',
    })


@login_required
@role_required(SUPERVISOR, MANAGER)
@require_POST
def shift_delete(request, pk):
    shift = get_object_or_404(permitted_shifts(request.user).select_related('unit'), pk=pk)
    if not can_manage_shift(request.user, shift):
        raise PermissionDenied('You cannot delete this shift count.')
    log_audit(request.user, 'shift_deleted', object_type='ShiftCount', object_id=shift.pk,
              target_label=str(shift), detail={'unit': shift.unit_id, 'date': shift.date.isoformat(),
                                               'shift': shift.shift, 'on_duty': shift.on_duty},
              request=request)
    shift.delete()
    messages.success(request, 'Shift count deleted.')
    return redirect('shifts')


# ---------------------------------------------------------------------------
# Notifications and audit trail
# ---------------------------------------------------------------------------
@login_required
def notifications_page(request):
    qs = Notification.objects.filter(user=request.user)
    f_status = request.GET.get('status') or ''
    if f_status == 'unread':
        qs = qs.filter(is_read=False)
    elif f_status == 'read':
        qs = qs.filter(is_read=True)
    page = _page(request, qs.order_by('-created_at'), 30)
    return render(request, 'core/notifications.html', {
        'notifications': page,
        'page_obj': page,
        'f_status': f_status,
        'unread': Notification.objects.filter(user=request.user, is_read=False).count(),
        'page_title': 'Notifications',
        'breadcrumb': 'Home / Notifications',
    })


@login_required
@require_POST
def notification_read(request, pk):
    note = get_object_or_404(Notification, pk=pk, user=request.user)
    if not note.is_read:
        note.is_read = True
        note.read_at = timezone.now()
        note.save(update_fields=['is_read', 'read_at'])
    next_url = request.POST.get('next') or request.META.get('HTTP_REFERER') or reverse('notifications')
    return redirect(next_url)


@login_required
@require_POST
def notifications_read_all(request):
    Notification.objects.filter(user=request.user, is_read=False).update(is_read=True, read_at=timezone.now())
    messages.success(request, 'All notifications marked as read.')
    return redirect('notifications')


@login_required
@role_required(SUPERVISOR)
def audit_page(request):
    qs = AuditEvent.objects.select_related('actor')
    q = (request.GET.get('q') or '').strip()
    event = request.GET.get('event') or ''
    object_type = request.GET.get('object_type') or ''
    actor_id = request.GET.get('actor') or ''
    start = _date_or_none(request.GET.get('start'))
    end = _date_or_none(request.GET.get('end'))
    if q:
        qs = qs.filter(
            Q(actor__username__icontains=q) | Q(event__icontains=q)
            | Q(object_type__icontains=q) | Q(target_label__icontains=q)
        )
    if event:
        qs = qs.filter(event=event)
    if object_type:
        qs = qs.filter(object_type=object_type)
    if actor_id.isdigit():
        qs = qs.filter(actor_id=int(actor_id))
    if start:
        qs = qs.filter(occurred_at__date__gte=start)
    if end:
        qs = qs.filter(occurred_at__date__lte=end)
    page = _page(request, qs.order_by('-occurred_at'), 50)
    return render(request, 'core/audit.html', {
        'events': page,
        'page_obj': page,
        'q': q, 'event': event, 'object_type': object_type, 'actor_id': actor_id,
        'start': start, 'end': end,
        'events_choices': AuditEvent.objects.values_list('event', flat=True).distinct().order_by('event'),
        'object_types': AuditEvent.objects.values_list('object_type', flat=True).distinct().order_by('object_type'),
        'actors': User.objects.filter(auditevent__isnull=False).distinct().order_by('username'),
        'page_title': 'Audit Trail',
        'breadcrumb': 'Home / Audit Trail',
    })
